from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, send_from_directory, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import sqlite3
import os
import uuid
import hashlib
import base64
import io
import json
import re
import unicodedata
from PIL import Image
from openpyxl import load_workbook
try:
    import qrcode
except ImportError:
    qrcode = None
from datetime import date, datetime
from functools import wraps
from typing import Any
from face_service import (
    FaceImageError,
    FaceRecognitionUnavailable,
    closest_match,
    extract_encodings,
    extract_single_encoding_from_file,
    image_data_url_to_bytes,
    serialize_encoding,
)

HORARIOS_INICIO = {'mañana': (8, 0), 'tarde': (13, 0)}

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STATIC_FOLDER = os.path.join(PROJECT_ROOT, 'static')
DATABASE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'database.db')
APP_FOLDER = os.path.dirname(os.path.abspath(__file__))
BACKUP_FOLDER = os.path.join(APP_FOLDER, 'backup_db')
LEGACY_BACKUP_FOLDER = os.path.join(APP_FOLDER, 'backups')
IMPORT_FOLDER = os.path.join(APP_FOLDER, '.import_staging')

app = Flask(__name__, static_folder=STATIC_FOLDER)
app.secret_key = 'clave_secreta_cambiar_en_produccion'
app.config['UPLOAD_FOLDER'] = os.path.join(STATIC_FOLDER, 'uploads')
app.config['ALLOWED_EXTENSIONS'] = {'png', 'jpg', 'jpeg', 'gif'}

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
os.makedirs(BACKUP_FOLDER, exist_ok=True)
os.makedirs(IMPORT_FOLDER, exist_ok=True)
app.config['MAX_CONTENT_LENGTH'] = 10 * 1024 * 1024
app.config['FACE_MATCH_TOLERANCE'] = float(os.environ.get('FACE_MATCH_TOLERANCE', '0.48'))

# Preserve existing backup files and their database history when moving to backup_db/.
if os.path.isdir(LEGACY_BACKUP_FOLDER):
    for legacy_name in os.listdir(LEGACY_BACKUP_FOLDER):
        legacy_path = os.path.join(LEGACY_BACKUP_FOLDER, legacy_name)
        if os.path.isfile(legacy_path) and os.path.splitext(legacy_name)[1].lower() in ('.db', '.sqlite', '.sqlite3'):
            destination = os.path.join(BACKUP_FOLDER, legacy_name)
            if not os.path.exists(destination):
                try:
                    os.replace(legacy_path, destination)
                except OSError:
                    app.logger.exception('No se pudo trasladar el backup heredado %s', legacy_name)
for staged_name in os.listdir(IMPORT_FOLDER):
    staged_path = os.path.join(IMPORT_FOLDER, staged_name)
    try:
        if os.path.isfile(staged_path) and os.path.getmtime(staged_path) < datetime.now().timestamp() - 86400:
            os.remove(staged_path)
    except OSError:
        app.logger.exception('No se pudo limpiar una planilla temporal')

# ---------- Base de datos ----------
def get_db():
    conn = sqlite3.connect(DATABASE_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    conn.execute('''
        CREATE TABLE IF NOT EXISTS usuarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            email TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL,
            nombre TEXT NOT NULL,
            apellido TEXT NOT NULL,
            dni TEXT,
            rol TEXT NOT NULL CHECK(rol IN ('directivo','preceptor','alumno','entrada')),
            curso TEXT,
            turno TEXT CHECK(turno IN ('mañana','tarde')),
            hora_entrada_oficial TEXT,
            capacitacion TEXT CHECK(capacitacion IN ('Informática','Electrónica')),
            profile_photo TEXT,
            qr_token TEXT,
            face_encoding TEXT,
            face_consent INTEGER NOT NULL DEFAULT 0
        )
    ''')
    conn.execute('''
        CREATE TABLE IF NOT EXISTS asistencias (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            alumno_id INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            estado TEXT NOT NULL CHECK(estado IN ('presente','ausente','tardanza')),
            hora TEXT,
            FOREIGN KEY (alumno_id) REFERENCES usuarios(id)
        )
    ''')
    conn.execute('''
        CREATE TABLE IF NOT EXISTS backups (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filename TEXT UNIQUE NOT NULL,
            tipo TEXT NOT NULL CHECK(tipo IN ('creado', 'cargado')),
            fecha TEXT NOT NULL,
            tamano INTEGER NOT NULL
        )
    ''')

    def add_column_if_missing(table, column, definition):
        existing = [row[1] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()]
        if column not in existing:
            conn.execute(f"ALTER TABLE {table} ADD COLUMN {definition}")

    add_column_if_missing('usuarios', 'apellido', 'apellido TEXT')
    add_column_if_missing('usuarios', 'dni', 'dni TEXT')
    add_column_if_missing('usuarios', 'turno', 'turno TEXT')
    add_column_if_missing('usuarios', 'hora_entrada_oficial', 'hora_entrada_oficial TEXT')
    add_column_if_missing('usuarios', 'capacitacion', 'capacitacion TEXT')
    add_column_if_missing('usuarios', 'profile_photo', 'profile_photo TEXT')
    add_column_if_missing('usuarios', 'qr_token', 'qr_token TEXT')
    add_column_if_missing('usuarios', 'face_encoding', 'face_encoding TEXT')
    add_column_if_missing('usuarios', 'face_consent', 'face_consent INTEGER NOT NULL DEFAULT 0')
    add_column_if_missing('asistencias', 'hora', 'hora TEXT')
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_usuarios_dni_unico ON usuarios(dni) WHERE dni IS NOT NULL AND dni != ''")

    usuarios_schema = conn.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='usuarios'").fetchone()[0]
    if "'entrada'" not in usuarios_schema:
        conn.execute('ALTER TABLE usuarios RENAME TO usuarios_antigua')
        conn.execute('''
            CREATE TABLE usuarios (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                email TEXT UNIQUE NOT NULL,
                password TEXT NOT NULL,
                nombre TEXT NOT NULL,
                apellido TEXT NOT NULL,
                dni TEXT,
                rol TEXT NOT NULL CHECK(rol IN ('directivo','preceptor','alumno','entrada')),
                curso TEXT,
                turno TEXT CHECK(turno IN ('mañana','tarde')),
                hora_entrada_oficial TEXT,
                capacitacion TEXT CHECK(capacitacion IN ('Informática','Electrónica')),
                profile_photo TEXT,
                qr_token TEXT,
                face_encoding TEXT,
                face_consent INTEGER NOT NULL DEFAULT 0
            )
        ''')
        conn.execute('''
            INSERT INTO usuarios (id, email, password, nombre, apellido, dni, rol, curso,
                                  turno, hora_entrada_oficial, capacitacion, profile_photo, qr_token, face_encoding, face_consent)
                 SELECT id, email, password, nombre, COALESCE(apellido, ''), dni, rol, curso, turno,
                     hora_entrada_oficial, capacitacion, profile_photo, qr_token, face_encoding, face_consent
            FROM usuarios_antigua
        ''')
        conn.execute('DROP TABLE usuarios_antigua')

    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_usuarios_dni_unico ON usuarios(dni) WHERE dni IS NOT NULL AND dni != ''")

    # Crear admin por defecto si no existe
    try:
        conn.execute("INSERT INTO usuarios (email, password, nombre, apellido, dni, rol, curso, turno, capacitacion) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                     ('admin@colegio.com', generate_password_hash('admin123'), 'Administrador', 'Admin', '', 'directivo', None, None, None))
    except sqlite3.IntegrityError:
        pass
    conn.commit()
    conn.close()

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']

def validate_dni(dni):
    dni = (dni or '').strip()
    if not dni or len(dni) > 8 or not dni.isdigit():
        raise ValueError('El DNI debe contener entre 1 y 8 números.')
    return dni


def validate_entry_time(turno: str | None, value: str | None = None) -> str | None:
    """Valida la hora oficial por usuario; deriva 08:00/13:00 si se omite."""
    if turno not in HORARIOS_INICIO:
        return None
    default_hour, default_minute = HORARIOS_INICIO[turno]
    candidate = (value or f'{default_hour:02d}:{default_minute:02d}').strip()
    try:
        return datetime.strptime(candidate, '%H:%M').strftime('%H:%M')
    except ValueError as error:
        raise ValueError('La hora oficial debe tener el formato HH:MM.') from error


def get_official_start(alumno) -> tuple[int, int] | None:
    """Obtiene el horario oficial de SQLite, con compatibilidad heredada."""
    official_time = alumno['hora_entrada_oficial'] if 'hora_entrada_oficial' in alumno.keys() else None
    if official_time:
        try:
            parsed_time = datetime.strptime(official_time, '%H:%M').time()
            return parsed_time.hour, parsed_time.minute
        except (TypeError, ValueError):
            app.logger.warning('Horario oficial inválido en la ficha del alumno %s', alumno['id'])
    return HORARIOS_INICIO.get(alumno['turno'])

IMPORT_FIELDS = [
    ('nombre', 'Nombre'),
    ('apellido', 'Apellido'),
    ('dni', 'DNI'),
    ('email', 'Correo electrónico'),
    ('password', 'Contraseña inicial'),
    ('curso', 'Curso'),
    ('turno', 'Turno'),
    ('capacitacion', 'Capacitación (opcional)'),
]
REQUIRED_IMPORT_FIELDS = {'nombre', 'apellido', 'dni', 'email', 'password', 'curso', 'turno'}
IMPORT_HEADER_ALIASES = {
    'nombre': ('nombre', 'nombres', 'name', 'first name'),
    'apellido': ('apellido', 'apellidos', 'surname', 'last name'),
    'dni': ('dni', 'documento', 'numero documento', 'nro documento', 'legajo dni'),
    'email': ('email', 'e mail', 'correo', 'correo electronico', 'mail'),
    'password': ('password', 'contrasena', 'clave', 'clave inicial', 'contrasena inicial'),
    'curso': ('curso', 'division', 'grado', 'curso division'),
    'turno': ('turno', 'jornada'),
    'capacitacion': ('capacitacion', 'especialidad'),
}

def normalize_import_header(value):
    normalized = unicodedata.normalize('NFKD', str(value).strip().casefold())
    without_accents = ''.join(character for character in normalized if not unicodedata.combining(character))
    return ' '.join(re.findall(r'[a-z0-9]+', without_accents))

def detect_import_mapping(headers):
    mapping = {field: None for field, _label in IMPORT_FIELDS}
    used_columns = set()
    normalized_headers = [normalize_import_header(header) for header in headers]
    for field, _label in IMPORT_FIELDS:
        aliases = [normalize_import_header(alias) for alias in IMPORT_HEADER_ALIASES[field]]
        matches = [
            index for index, header in enumerate(normalized_headers)
            if index not in used_columns and header and any(
                header == alias or (len(alias) >= 5 and alias in header)
                for alias in aliases
            )
        ]
        if matches:
            mapping[field] = matches[0]
            used_columns.add(matches[0])
    return mapping

def excel_value(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.isoformat(sep=' ', timespec='minutes')
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

def bulk_import_path(import_id):
    if not import_id or not re.fullmatch(r'[a-f0-9]{32}', import_id):
        return None
    return os.path.join(IMPORT_FOLDER, f'{import_id}.json')

def load_bulk_import(import_id):
    path = bulk_import_path(import_id)
    if not path or not os.path.isfile(path):
        return None
    try:
        with open(path, encoding='utf-8') as file:
            return json.load(file)
    except (OSError, json.JSONDecodeError):
        return None

def save_bulk_import(import_id, payload):
    path = bulk_import_path(import_id)
    if not path:
        raise ValueError('Identificador de importación inválido.')
    temporary_path = f'{path}.tmp'
    with open(temporary_path, 'w', encoding='utf-8') as file:
        json.dump(payload, file, ensure_ascii=False)
    os.replace(temporary_path, path)

def delete_bulk_import(import_id):
    path = bulk_import_path(import_id)
    if path and os.path.isfile(path):
        os.remove(path)

def validate_import_mapping(mapping, headers):
    errors = []
    for field in REQUIRED_IMPORT_FIELDS:
        if mapping.get(field) is None:
            label = dict(IMPORT_FIELDS)[field]
            errors.append(f'Asigná una columna para {label}.')
    selected = [column for column in mapping.values() if column is not None]
    if len(selected) != len(set(selected)):
        errors.append('Cada campo debe usar una columna diferente.')
    if any(not isinstance(column, int) or column < 0 or column >= len(headers) for column in selected):
        errors.append('La selección contiene una columna inválida.')
    return errors

def map_import_rows(rows, mapping):
    records = []
    for row_number, row in enumerate(rows, start=2):
        record = {'_row': row_number}
        for field, _label in IMPORT_FIELDS:
            index = mapping.get(field)
            record[field] = row[index].strip() if index is not None and index < len(row) else ''
        records.append(record)
    return records

def validate_student_records(records, conn=None):
    own_connection = conn is None
    connection = conn or get_db()
    row_errors = []
    seen_dns = set()
    seen_emails = set()
    try:
        for record in records:
            errors = []
            row_number = record.get('_row', '?')
            name = record.get('nombre', '').strip()
            surname = record.get('apellido', '').strip()
            dni = record.get('dni', '').strip()
            email = record.get('email', '').strip()
            password = record.get('password', '').strip()
            course = record.get('curso', '').strip().upper()
            shift = record.get('turno', '').strip().lower()
            training = record.get('capacitacion', '').strip()
            if not name:
                errors.append('falta el nombre')
            if not surname:
                errors.append('falta el apellido')
            try:
                validate_dni(dni)
            except ValueError:
                errors.append('DNI inválido (debe tener entre 1 y 8 dígitos)')
            if dni:
                if dni in seen_dns:
                    errors.append('DNI repetido en la planilla')
                elif connection.execute('SELECT 1 FROM usuarios WHERE dni=?', (dni,)).fetchone():
                    errors.append('DNI ya registrado')
                seen_dns.add(dni)
            if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', email):
                errors.append('correo electrónico inválido')
            if email:
                normalized_email = email.casefold()
                if normalized_email in seen_emails:
                    errors.append('correo repetido en la planilla')
                elif connection.execute('SELECT 1 FROM usuarios WHERE lower(email)=lower(?)', (email,)).fetchone():
                    errors.append('correo ya registrado')
                seen_emails.add(normalized_email)
            if not password:
                errors.append('falta la contraseña inicial')
            if not re.fullmatch(r'[1-6][AB]', course):
                errors.append('curso inválido (ejemplo: 2A)')
            if shift not in ('mañana', 'tarde'):
                errors.append('turno inválido (mañana o tarde)')
            if training and training not in ('Informática', 'Electrónica'):
                errors.append('capacitación inválida')
            record['curso'] = course
            record['turno'] = shift
            if re.fullmatch(r'[1-6][AB]', course) and int(course[0]) < 3:
                record['capacitacion'] = ''
            if errors:
                row_errors.append({'row': row_number, 'messages': errors})
        return row_errors
    finally:
        if own_connection:
            connection.close()

def valid_backup_file(path):
    try:
        connection = sqlite3.connect(f'file:{path}?mode=ro', uri=True)
        result = connection.execute('PRAGMA integrity_check').fetchone()[0]
        connection.close()
        return result == 'ok'
    except sqlite3.Error:
        return False

def register_backup(filename, backup_type):
    conn = get_db()
    conn.execute('INSERT OR REPLACE INTO backups (filename, tipo, fecha, tamano) VALUES (?, ?, ?, ?)',
                 (filename, backup_type, datetime.now().isoformat(timespec='seconds'),
                  os.path.getsize(os.path.join(BACKUP_FOLDER, filename))))
    conn.commit()
    conn.close()

def save_photo(photo_file=None, photo_data=None):
    if photo_file and photo_file.filename:
        if not allowed_file(photo_file.filename):
            raise ValueError('Formato de foto no permitido. Usa png, jpg, jpeg o gif.')
        filename = f"{uuid.uuid4().hex}_{secure_filename(photo_file.filename)}"
        photo_file.save(os.path.join(app.config['UPLOAD_FOLDER'], filename))
        return filename
    if photo_data:
        if not photo_data.startswith('data:image/'):
            raise ValueError('La captura de cámara no es una imagen válida.')
        try:
            encoded = photo_data.split(',', 1)[1]
            image = Image.open(io.BytesIO(base64.b64decode(encoded)))
            image = image.convert('RGB')
            filename = f'{uuid.uuid4().hex}.jpg'
            image.save(os.path.join(app.config['UPLOAD_FOLDER'], filename), format='JPEG', quality=90)
            return filename
        except (ValueError, IndexError, base64.binascii.Error, OSError) as error:
            raise ValueError('No se pudo guardar la foto capturada.') from error
    return None

def generate_seed_word(dni: str, user_id: int) -> str:
    dice = f"{dni}|{user_id}"
    digest = hashlib.sha256(dice.encode('utf-8')).digest()
    return base64.urlsafe_b64encode(digest[:12]).decode('utf-8').rstrip('=')

def generate_qr_token(dni: str, user_id: int) -> str:
    seed = generate_seed_word(dni, user_id)
    digest = hashlib.pbkdf2_hmac('sha256', seed.encode('utf-8'), app.secret_key.encode('utf-8'), 100000, dklen=18)
    return base64.urlsafe_b64encode(digest).decode('utf-8').rstrip('=')

def make_qr_base64(data: str) -> str:
    buffer = io.BytesIO()
    make_qr_image(data).save(buffer, format='PNG')
    buffer.seek(0)
    return base64.b64encode(buffer.read()).decode('utf-8')

def make_qr_image(data: str):
    if qrcode is None:
        raise RuntimeError('La librería qrcode no está instalada. Instala qrcode[pil] para generar códigos QR.')
    qr = qrcode.QRCode(box_size=10, border=3)
    qr.add_data(data)
    qr.make(fit=True)
    return qr.make_image(fill_color='black', back_color='white').convert('RGB')

def update_user_qr_token(conn, user_id: int, dni: str) -> str | None:
    qr_token = None
    if dni:
        qr_token = generate_qr_token(dni, user_id)
    conn.execute("UPDATE usuarios SET qr_token=? WHERE id=?", (qr_token, user_id))
    return qr_token

init_db()

# ---------- Decoradores de autenticación ----------
def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'usuario_id' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated

def rol_requerido(*roles):
    def decorator(f):
        @wraps(f)
        def decorated(*args, **kwargs):
            if session.get('rol') not in roles:
                flash('No tienes permiso para acceder a esta página.')
                return redirect(url_for('dashboard'))
            return f(*args, **kwargs)
        return decorated
    return decorator

# ---------- Rutas ----------
@app.route('/')
def index():
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        email = request.form['email']
        password = request.form['password']
        conn = get_db()
        user = conn.execute("SELECT * FROM usuarios WHERE email = ?", (email,)).fetchone()
        conn.close()
        if user and check_password_hash(user['password'], password):
            session['usuario_id'] = user['id']
            session['nombre'] = user['nombre']
            session['rol'] = user['rol']
            session['curso'] = user['curso']  # puede ser None
            return redirect(url_for('dashboard'))
        flash('Correo o contraseña incorrectos.')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/dashboard')
@login_required
def dashboard():
    rol = session['rol']
    if rol == 'directivo':
        return redirect(url_for('dashboard_directivo'))
    elif rol == 'preceptor':
        return redirect(url_for('dashboard_preceptor'))
    elif rol == 'alumno':
        return redirect(url_for('dashboard_alumno'))
    elif rol == 'entrada':
        return redirect(url_for('dashboard_entrada'))
    return redirect(url_for('login'))

@app.route('/directivo')
@login_required
@rol_requerido('directivo')
def dashboard_directivo():
    conn = get_db()
    total_alumnos = conn.execute("SELECT COUNT(*) FROM usuarios WHERE rol='alumno'").fetchone()[0]
    total_preceptores = conn.execute("SELECT COUNT(*) FROM usuarios WHERE rol='preceptor'").fetchone()[0]
    conn.close()
    return render_template('dashboard_directivo.html', total_alumnos=total_alumnos,
                           total_preceptores=total_preceptores)

@app.route('/preceptor')
@login_required
@rol_requerido('preceptor')
def dashboard_preceptor():
    return render_template('dashboard_preceptor.html', curso=session.get('curso'))

@app.route('/alumno')
@login_required
@rol_requerido('alumno')
def dashboard_alumno():
    return render_template('dashboard_alumno.html')

@app.route('/entrada')
@login_required
@rol_requerido('entrada')
def dashboard_entrada():
    return render_template('dashboard_entrada.html')

# --- Gestión de usuarios (solo directivo) ---
@app.route('/usuarios')
@login_required
@rol_requerido('directivo')
def gestion_usuarios():
    search = request.args.get('q', '').strip()
    conn = get_db()
    if search:
        pattern = f'%{search}%'
        usuarios = conn.execute(
            "SELECT * FROM usuarios WHERE dni LIKE ? OR nombre LIKE ? OR apellido LIKE ? ORDER BY rol, nombre",
            (pattern, pattern, pattern)
        ).fetchall()
    else:
        usuarios = conn.execute("SELECT * FROM usuarios ORDER BY rol, nombre").fetchall()
    conn.close()
    return render_template('gestion_usuarios.html', usuarios=usuarios, search=search)

@app.route('/carga-masiva', methods=['GET', 'POST'])
@login_required
@rol_requerido('directivo')
def carga_masiva():
    import_id = session.get('bulk_import_id')
    payload = load_bulk_import(import_id) if import_id else None
    if request.method == 'GET':
        return render_template('carga_masiva.html', stage='upload', payload=payload)

    action = request.form.get('action')
    if action == 'upload':
        if import_id:
            delete_bulk_import(import_id)
            session.pop('bulk_import_id', None)
        workbook_file = request.files.get('excel')
        if not workbook_file or not workbook_file.filename or not workbook_file.filename.lower().endswith('.xlsx'):
            flash('Seleccioná una planilla Excel .xlsx.')
            return redirect(url_for('carga_masiva'))
        try:
            workbook = load_workbook(io.BytesIO(workbook_file.read()), read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            first_row = next(iterator, None)
            if not first_row:
                raise ValueError('La planilla está vacía.')
            headers = [str(value).strip() if value is not None else '' for value in first_row]
            headers = [header or f'Columna {index + 1}' for index, header in enumerate(headers)]
            rows = []
            for row_number, row in enumerate(iterator, start=2):
                if row_number > 10001:
                    raise ValueError('La planilla supera el límite de 10.000 filas.')
                values = [excel_value(value) for value in row]
                values.extend([''] * max(0, len(headers) - len(values)))
                values = values[:len(headers)]
                if any(value.strip() for value in values):
                    rows.append(values)
            workbook.close()
            if not rows:
                raise ValueError('La planilla contiene encabezados, pero no tiene filas de alumnos.')
            if len(headers) > 100:
                raise ValueError('La planilla supera el límite de 100 columnas.')
            import_id = uuid.uuid4().hex
            mapping = detect_import_mapping(headers)
            payload = {'headers': headers, 'rows': rows, 'filename': secure_filename(workbook_file.filename),
                       'mapping': mapping}
            mapping_errors = validate_import_mapping(mapping, headers)
            if not mapping_errors:
                records = map_import_rows(rows, mapping)
                payload.update({
                    'records': records,
                    'row_errors': validate_student_records(records),
                })
                stage = 'preview'
            else:
                stage = 'mapping'
            save_bulk_import(import_id, payload)
            session['bulk_import_id'] = import_id
            return render_template('carga_masiva.html', stage=stage, payload=payload,
                                   fields=IMPORT_FIELDS, mapping=mapping, errors=mapping_errors)
        except (OSError, ValueError, KeyError, StopIteration) as error:
            flash(f'No se pudo leer la planilla: {error}')
            return redirect(url_for('carga_masiva'))
        except Exception as error:
            app.logger.exception('Error leyendo la planilla Excel')
            flash(f'No se pudo leer la planilla Excel: {error}')
            return redirect(url_for('carga_masiva'))

    if not import_id or not payload:
        flash('La vista previa expiró. Volvé a seleccionar la planilla.')
        session.pop('bulk_import_id', None)
        return redirect(url_for('carga_masiva'))

    if action == 'map':
        mapping = {}
        for field, _label in IMPORT_FIELDS:
            selected = request.form.get(f'map_{field}', '')
            mapping[field] = int(selected) if selected.isdigit() else None
        errors = validate_import_mapping(mapping, payload['headers'])
        if errors:
            payload['mapping'] = mapping
            save_bulk_import(import_id, payload)
            return render_template('carga_masiva.html', stage='mapping', payload=payload,
                                   fields=IMPORT_FIELDS, mapping=mapping, errors=errors)
        records = map_import_rows(payload['rows'], mapping)
        row_errors = validate_student_records(records)
        payload.update({'mapping': mapping, 'records': records, 'row_errors': row_errors})
        save_bulk_import(import_id, payload)
        return render_template('carga_masiva.html', stage='preview', payload=payload,
                               fields=IMPORT_FIELDS, mapping=mapping, errors=[])

    if action == 'import':
        records = payload.get('records', [])
        if not records:
            flash('Primero asociá las columnas y validá la vista previa.')
            return render_template('carga_masiva.html', stage='mapping', payload=payload,
                                   fields=IMPORT_FIELDS, mapping=payload.get('mapping', {}), errors=[])
        row_errors = validate_student_records(records)
        if row_errors:
            payload['row_errors'] = row_errors
            save_bulk_import(import_id, payload)
            return render_template('carga_masiva.html', stage='preview', payload=payload,
                                   fields=IMPORT_FIELDS, mapping=payload.get('mapping', {}), errors=[])
        try:
            # Backup is mandatory and must succeed before inserting any student.
            backup_name = create_database_backup(backup_type='creado')
            conn = get_db()
            try:
                conn.execute('BEGIN IMMEDIATE')
                concurrent_errors = validate_student_records(records, conn=conn)
                if concurrent_errors:
                    conn.rollback()
                    payload['row_errors'] = concurrent_errors
                    save_bulk_import(import_id, payload)
                    return render_template('carga_masiva.html', stage='preview', payload=payload,
                                           fields=IMPORT_FIELDS, mapping=payload.get('mapping', {}), errors=[],
                                           backup_created=backup_name)
                for record in records:
                    cursor = conn.execute(
                        "INSERT INTO usuarios (email, password, nombre, apellido, dni, rol, curso, turno, "
                        "hora_entrada_oficial, capacitacion) VALUES (?, ?, ?, ?, ?, 'alumno', ?, ?, ?, ?)",
                        (record['email'], generate_password_hash(record['password']), record['nombre'],
                         record['apellido'], record['dni'], record['curso'], record['turno'],
                         validate_entry_time(record['turno']), record['capacitacion'] or None)
                    )
                    update_user_qr_token(conn, cursor.lastrowid, record['dni'])
                conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                conn.close()
        except Exception:
            app.logger.exception('Falló la importación masiva; se revirtió la transacción')
            flash('No se pudo completar la importación. No se aplicaron cambios; revisá el backup y los datos.')
            return render_template('carga_masiva.html', stage='preview', payload=payload,
                                   fields=IMPORT_FIELDS, mapping=payload.get('mapping', {}), errors=[])

        imported = [{'nombre': row['nombre'], 'apellido': row['apellido'], 'dni': row['dni'], 'email': row['email']}
                    for row in records]
        delete_bulk_import(import_id)
        session.pop('bulk_import_id', None)
        return render_template('carga_masiva.html', stage='done', imported=imported,
                               backup_name=backup_name)

    flash('Acción de carga no reconocida.')
    return redirect(url_for('carga_masiva'))

def create_database_backup(backup_type='creado'):
    filename = f"backup_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.db"
    path = os.path.join(BACKUP_FOLDER, filename)
    source = sqlite3.connect(DATABASE_PATH)
    destination = sqlite3.connect(path)
    try:
        source.backup(destination)
    finally:
        destination.close()
        source.close()
    register_backup(filename, backup_type)
    return filename

def restore_database_backup(backup_path):
    source = sqlite3.connect(backup_path)
    destination = sqlite3.connect(DATABASE_PATH)
    try:
        source.backup(destination)
        destination.commit()
    finally:
        destination.close()
        source.close()

@app.route('/backups', methods=['GET', 'POST'])
@login_required
@rol_requerido('directivo')
def gestion_backups():
    if request.method == 'POST':
        action = request.form.get('action')
        if action == 'crear':
            create_database_backup()
            flash('Backup creado correctamente.')
        elif action == 'cargar':
            backup_file = request.files.get('backup')
            if not backup_file or not backup_file.filename:
                flash('Seleccioná un archivo de backup.')
            else:
                original_name = secure_filename(backup_file.filename)
                extension = os.path.splitext(original_name)[1].lower()
                if extension not in ('.db', '.sqlite', '.sqlite3'):
                    flash('El backup debe tener extensión .db, .sqlite o .sqlite3.')
                else:
                    uploaded_name = f"cargado_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}{extension}"
                    uploaded_path = os.path.join(BACKUP_FOLDER, uploaded_name)
                    backup_file.save(uploaded_path)
                    if not valid_backup_file(uploaded_path):
                        os.remove(uploaded_path)
                        flash('El archivo no es una base SQLite válida.')
                    else:
                        try:
                            create_database_backup()
                            restore_database_backup(uploaded_path)
                            init_db()
                            register_backup(uploaded_name, 'cargado')
                            flash('Backup cargado y restaurado correctamente.')
                        except (sqlite3.Error, OSError):
                            flash('No se pudo restaurar el backup. La base actual no fue reemplazada.')
        return redirect(url_for('gestion_backups'))

    conn = get_db()
    backups = conn.execute('SELECT filename, tipo, fecha, tamano FROM backups ORDER BY fecha DESC').fetchall()
    conn.close()
    backups = [backup for backup in backups if os.path.isfile(os.path.join(BACKUP_FOLDER, backup['filename']))]
    return render_template('backups.html', backups=backups)

@app.route('/backups/descargar/<path:filename>')
@login_required
@rol_requerido('directivo')
def descargar_backup(filename):
    safe_filename = os.path.basename(filename)
    return send_from_directory(BACKUP_FOLDER, safe_filename, as_attachment=True)

@app.route('/usuarios/nuevo', methods=['GET', 'POST'])
@login_required
@rol_requerido('directivo')
def nuevo_usuario():
    if request.method == 'POST':
        email = request.form['email']
        password = generate_password_hash(request.form['password'])
        nombre = request.form['nombre']
        apellido = request.form['apellido']
        try:
            dni = validate_dni(request.form.get('dni'))
        except ValueError as error:
            flash(str(error))
            return render_template('editar_usuario.html', usuario=None, titulo='Nuevo Usuario')
        rol = request.form['rol']
        curso = request.form.get('curso') if rol in ('preceptor', 'alumno') else None
        turno = request.form.get('turno') if rol in ('preceptor', 'alumno') else None
        try:
            hora_entrada_oficial = validate_entry_time(turno, request.form.get('hora_entrada_oficial'))
        except ValueError as error:
            flash(str(error))
            return render_template('editar_usuario.html', usuario=None, titulo='Nuevo Usuario')
        capacitacion = request.form.get('capacitacion') if rol == 'alumno' and curso and curso[0].isdigit() and int(curso[0]) >= 3 else None
        photo_filename = None
        photo_file = request.files.get('foto')
        try:
            photo_filename = save_photo(photo_file, request.form.get('foto_data'))
        except ValueError as error:
            flash(str(error))
            return render_template('editar_usuario.html', usuario=None, titulo='Nuevo Usuario')
        face_consent = int(rol == 'alumno' and request.form.get('face_consent') == '1')
        face_encoding = None
        if face_consent and not photo_filename:
            flash('Para registrar reconocimiento facial, adjuntá una foto del alumno además de otorgar consentimiento.')
            return render_template('editar_usuario.html', usuario=None, titulo='Nuevo Usuario')
        if face_consent and photo_filename:
            try:
                face_encoding = serialize_encoding(extract_single_encoding_from_file(
                    os.path.join(app.config['UPLOAD_FOLDER'], photo_filename)
                ))
            except (FaceImageError, FaceRecognitionUnavailable) as error:
                try:
                    os.remove(os.path.join(app.config['UPLOAD_FOLDER'], photo_filename))
                except OSError:
                    app.logger.exception('No se pudo limpiar una foto de registro facial inválida')
                flash(str(error))
                return render_template('editar_usuario.html', usuario=None, titulo='Nuevo Usuario')
        conn = get_db()
        try:
            cursor = conn.execute("INSERT INTO usuarios (email, password, nombre, apellido, dni, rol, curso, turno, hora_entrada_oficial, capacitacion, profile_photo, face_encoding, face_consent) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                         (email, password, nombre, apellido, dni, rol, curso, turno,
                          hora_entrada_oficial, capacitacion, photo_filename, face_encoding, face_consent))
            user_id = cursor.lastrowid
            qr_token = update_user_qr_token(conn, user_id, dni)
            conn.commit()
        except sqlite3.IntegrityError as error:
            if photo_filename:
                try:
                    os.remove(os.path.join(app.config['UPLOAD_FOLDER'], photo_filename))
                except OSError:
                    app.logger.exception('No se pudo limpiar una foto de usuario duplicado')
            flash('El DNI ya está registrado.' if 'dni' in str(error).lower() else 'El correo ya está registrado.')
            conn.close()
            return render_template('editar_usuario.html', usuario=None, titulo='Nuevo Usuario')
        conn.close()
        flash('Usuario creado exitosamente.')
        return redirect(url_for('gestion_usuarios'))
    return render_template('editar_usuario.html', usuario=None, titulo='Nuevo Usuario')

@app.route('/usuarios/editar/<int:id>', methods=['GET', 'POST'])
@login_required
@rol_requerido('directivo')
def editar_usuario(id):
    conn = get_db()
    usuario = conn.execute("SELECT * FROM usuarios WHERE id = ?", (id,)).fetchone()
    if not usuario:
        flash('Usuario no encontrado.')
        conn.close()
        return redirect(url_for('gestion_usuarios'))
    if request.method == 'POST':
        nombre = request.form['nombre']
        apellido = request.form['apellido']
        try:
            dni = validate_dni(request.form.get('dni'))
        except ValueError as error:
            flash(str(error))
            conn.close()
            return render_template('editar_usuario.html', usuario=usuario, titulo='Editar Usuario')
        duplicate = conn.execute('SELECT id FROM usuarios WHERE dni=? AND id!=?', (dni, id)).fetchone()
        if duplicate:
            flash('El DNI ya está registrado por otro usuario.')
            conn.close()
            return render_template('editar_usuario.html', usuario=usuario, titulo='Editar Usuario')
        rol = request.form['rol']
        curso = request.form.get('curso') if rol in ('preceptor', 'alumno') else None
        turno = request.form.get('turno') if rol in ('preceptor', 'alumno') else None
        try:
            hora_entrada_oficial = validate_entry_time(turno, request.form.get('hora_entrada_oficial'))
        except ValueError as error:
            flash(str(error))
            conn.close()
            return render_template('editar_usuario.html', usuario=usuario, titulo='Editar Usuario')
        capacitacion = request.form.get('capacitacion') if rol == 'alumno' and curso and curso[0].isdigit() and int(curso[0]) >= 3 else None
        photo_filename = usuario['profile_photo']
        photo_file = request.files.get('foto')
        try:
            new_photo = save_photo(photo_file, request.form.get('foto_data'))
            if new_photo:
                photo_filename = new_photo
        except ValueError as error:
            flash(str(error))
            conn.close()
            return render_template('editar_usuario.html', usuario=usuario, titulo='Editar Usuario')
        face_consent = int(rol == 'alumno' and request.form.get('face_consent') == '1')
        face_encoding = usuario['face_encoding']
        if not face_consent:
            face_encoding = None
        elif new_photo or not face_encoding:
            face_encoding = None
            if photo_filename:
                try:
                    face_encoding = serialize_encoding(extract_single_encoding_from_file(
                        os.path.join(app.config['UPLOAD_FOLDER'], photo_filename)
                    ))
                except (FaceImageError, FaceRecognitionUnavailable) as error:
                    if new_photo:
                        try:
                            os.remove(os.path.join(app.config['UPLOAD_FOLDER'], new_photo))
                        except OSError:
                            app.logger.exception('No se pudo limpiar una foto facial inválida')
                    flash(str(error))
                    conn.close()
                    return render_template('editar_usuario.html', usuario=usuario, titulo='Editar Usuario')
            else:
                flash('Para registrar reconocimiento facial, adjuntá una foto del alumno además de otorgar consentimiento.')
                conn.close()
                return render_template('editar_usuario.html', usuario=usuario, titulo='Editar Usuario')
        qr_token = generate_qr_token(dni, id) if dni else None
        conn.execute("UPDATE usuarios SET nombre=?, apellido=?, dni=?, rol=?, curso=?, turno=?, hora_entrada_oficial=?, capacitacion=?, profile_photo=?, qr_token=?, face_encoding=?, face_consent=? WHERE id=?",
                 (nombre, apellido, dni, rol, curso, turno, hora_entrada_oficial,
                  capacitacion, photo_filename, qr_token, face_encoding, face_consent, id))
        if request.form['password']:
            conn.execute("UPDATE usuarios SET password=? WHERE id=?",
                         (generate_password_hash(request.form['password']), id))
        conn.commit()
        conn.close()
        flash('Usuario actualizado.')
        return redirect(url_for('gestion_usuarios'))
    conn.close()
    return render_template('editar_usuario.html', usuario=usuario, titulo='Editar Usuario')

@app.route('/usuarios/eliminar/<int:id>')
@login_required
@rol_requerido('directivo')
def eliminar_usuario(id):
    conn = get_db()
    conn.execute("DELETE FROM asistencias WHERE alumno_id=?", (id,))
    conn.execute("DELETE FROM usuarios WHERE id=? AND rol!='directivo'", (id,))
    conn.commit()
    conn.close()
    flash('Usuario eliminado.')
    return redirect(url_for('gestion_usuarios'))

@app.route('/usuarios/qrcode/<int:id>')
@login_required
@rol_requerido('directivo')
def ver_qr_usuario(id):
    conn = get_db()
    usuario = conn.execute("SELECT id, nombre, apellido, dni, qr_token FROM usuarios WHERE id = ?", (id,)).fetchone()
    if not usuario:
        flash('Usuario no encontrado.')
        conn.close()
        return redirect(url_for('gestion_usuarios'))
    qr_token = usuario['qr_token']
    if not qr_token and usuario['dni']:
        qr_token = generate_qr_token(usuario['dni'], usuario['id'])
        conn.execute("UPDATE usuarios SET qr_token=? WHERE id=?", (qr_token, usuario['id']))
        conn.commit()
    conn.close()
    if not qr_token:
        flash('No es posible generar QR sin DNI.')
        return redirect(url_for('gestion_usuarios'))
    # El QR contiene únicamente el token seguro (no incluye IP ni URL)
    qr_image = make_qr_base64(qr_token)
    return render_template('qr_usuario.html', usuario=usuario, qr_image=qr_image, token=qr_token)

@app.route('/usuarios/qrcode/<int:id>/descargar/<formato>')
@login_required
@rol_requerido('directivo')
def descargar_qr_usuario(id, formato):
    if formato not in ('jpg', 'pdf'):
        return 'Formato no soportado', 400
    conn = get_db()
    usuario = conn.execute(
        "SELECT id, nombre, apellido, dni, qr_token FROM usuarios WHERE id = ?", (id,)
    ).fetchone()
    if not usuario:
        conn.close()
        flash('Usuario no encontrado.')
        return redirect(url_for('gestion_usuarios'))
    qr_token = usuario['qr_token']
    if not qr_token and usuario['dni']:
        qr_token = generate_qr_token(usuario['dni'], usuario['id'])
        conn.execute("UPDATE usuarios SET qr_token=? WHERE id=?", (qr_token, usuario['id']))
        conn.commit()
    conn.close()
    if not qr_token:
        flash('No es posible generar QR sin DNI.')
        return redirect(url_for('gestion_usuarios'))

    image = make_qr_image(qr_token)
    buffer = io.BytesIO()
    if formato == 'pdf':
        image.save(buffer, format='PDF', resolution=100.0)
        mimetype = 'application/pdf'
    else:
        image.save(buffer, format='JPEG', quality=95)
        mimetype = 'image/jpeg'
    buffer.seek(0)
    filename = f"qr_{usuario['dni'] or usuario['id']}.{formato}"
    return send_file(buffer, mimetype=mimetype, as_attachment=True, download_name=filename)

@app.route('/qr/<token>')
@login_required
@rol_requerido('directivo', 'preceptor', 'entrada')
def ver_usuario_por_qr(token):
    conn = get_db()
    usuario = conn.execute("SELECT id, nombre, apellido, rol, curso, turno, capacitacion, profile_photo FROM usuarios WHERE qr_token = ?", (token,)).fetchone()
    asistencia_hoy = None
    if usuario and usuario['rol'] == 'alumno':
        asistencia_hoy = conn.execute(
            "SELECT hora, estado FROM asistencias WHERE alumno_id=? AND fecha=? ORDER BY id LIMIT 1",
            (usuario['id'], date.today().isoformat())
        ).fetchone()
    conn.close()
    if not usuario:
        flash('QR inválido o usuario no encontrado.')
        return redirect(url_for('dashboard'))
    return render_template('qr_validado.html', usuario=usuario, asistencia_hoy=asistencia_hoy)


@app.route('/qr/validate', methods=['POST'])
@login_required
@rol_requerido('directivo', 'preceptor', 'entrada')
def validate_qr():
    data = request.get_json(silent=True) or {}
    token = data.get('token')
    if not token:
        return jsonify({'ok': False, 'error': 'Token faltante'}), 400
    conn = get_db()
    usuario = conn.execute("SELECT id, nombre, apellido, rol, curso, turno, capacitacion, profile_photo FROM usuarios WHERE qr_token = ?", (token,)).fetchone()
    conn.close()
    if not usuario:
        return jsonify({'ok': False, 'error': 'Token inválido'}), 404
    return jsonify({'ok': True, 'usuario': {'id': usuario['id'], 'nombre': usuario['nombre'], 'apellido': usuario['apellido'], 'rol': usuario['rol'], 'curso': usuario['curso'], 'turno': usuario['turno'], 'capacitacion': usuario['capacitacion'], 'profile_photo': usuario['profile_photo']}})


def register_face_attendance(alumno_id: int) -> dict[str, Any]:
    """Registra una única entrada diaria y aplica un cooldown de cinco minutos."""
    ahora = datetime.now()
    fecha = ahora.date().isoformat()
    hora = ahora.strftime('%H:%M:%S')
    conn = get_db()
    try:
        conn.execute('BEGIN IMMEDIATE')
        alumno = conn.execute(
            "SELECT id, turno, hora_entrada_oficial FROM usuarios WHERE id=? AND rol='alumno' AND face_consent=1",
            (alumno_id,),
        ).fetchone()
        if not alumno:
            conn.rollback()
            return {'status': 'unavailable', 'message': 'El usuario no está habilitado para reconocimiento facial.'}

        existente = conn.execute(
            'SELECT hora, estado FROM asistencias WHERE alumno_id=? AND fecha=? ORDER BY id LIMIT 1',
            (alumno_id, fecha),
        ).fetchone()
        if existente:
            cooldown = 0
            try:
                hora_previa = datetime.strptime(existente['hora'], '%H:%M:%S').time()
                marcada = datetime.combine(ahora.date(), hora_previa)
                cooldown = max(0, 300 - int((ahora - marcada).total_seconds()))
            except (TypeError, ValueError):
                pass
            conn.commit()
            if cooldown:
                message = f"La entrada ya fue registrada hace menos de cinco minutos ({existente['hora'] or 'hora no disponible'})."
            else:
                message = f"La entrada de hoy ya fue registrada a las {existente['hora'] or 'hora no disponible'}."
            return {
                'status': 'duplicate',
                'message': message,
                'hora': existente['hora'],
                'estado': existente['estado'],
                'cooldown_seconds': cooldown,
            }

        inicio = get_official_start(alumno)
        estado = 'tardanza' if inicio and (ahora.hour, ahora.minute) > inicio else 'presente'
        conn.execute(
            'INSERT INTO asistencias (alumno_id, fecha, estado, hora) VALUES (?,?,?,?)',
            (alumno_id, fecha, estado, hora),
        )
        conn.commit()
        return {
            'status': 'registered',
            'message': f"Entrada registrada a las {hora}: {'Llegada tarde' if estado == 'tardanza' else 'A tiempo'}.",
            'hora': hora,
            'estado': estado,
            'cooldown_seconds': 300,
        }
    except sqlite3.Error:
        conn.rollback()
        app.logger.exception('No se pudo registrar asistencia facial del alumno %s', alumno_id)
        raise
    finally:
        conn.close()


@app.route('/asistencia/rostro/reconocer', methods=['POST'])
@login_required
@rol_requerido('entrada', 'preceptor')
def reconocer_rostro():
    """Analiza una captura de cámara y registra al alumno que coincida."""
    data = request.get_json(silent=True) or {}
    try:
        image_bytes = image_data_url_to_bytes(data.get('frame'))
        encodings = extract_encodings(image_bytes)
    except FaceRecognitionUnavailable as error:
        return jsonify({'ok': False, 'error': str(error), 'feature_unavailable': True}), 503
    except FaceImageError as error:
        return jsonify({'ok': False, 'error': str(error)}), 400

    if not encodings:
        return jsonify({'ok': True, 'face_detected': False, 'identified': False})
    if len(encodings) != 1:
        return jsonify({
            'ok': True,
            'face_detected': True,
            'identified': False,
            'multiple_faces': True,
            'message': 'Debe aparecer una sola persona frente a la cámara.',
        })

    conn = get_db()
    try:
        candidates = conn.execute(
            "SELECT id, nombre, apellido, turno, face_encoding FROM usuarios "
            "WHERE rol='alumno' AND face_consent=1 AND face_encoding IS NOT NULL"
        ).fetchall()
    finally:
        conn.close()

    try:
        match = closest_match(encodings[0], candidates, app.config['FACE_MATCH_TOLERANCE'])
    except FaceRecognitionUnavailable as error:
        return jsonify({'ok': False, 'error': str(error), 'feature_unavailable': True}), 503
    except FaceImageError as error:
        return jsonify({'ok': False, 'error': str(error)}), 400
    if not match:
        return jsonify({'ok': True, 'face_detected': True, 'identified': False})

    alumno, distance = match
    if data.get('confirm_user_id') != alumno['id']:
        return jsonify({
            'ok': True,
            'face_detected': True,
            'identified': True,
            'confirmed': False,
            'user': {'id': alumno['id'], 'nombre': alumno['nombre'], 'apellido': alumno['apellido']},
            'distance': round(distance, 4),
        })
    try:
        attendance = register_face_attendance(alumno['id'])
    except sqlite3.Error:
        return jsonify({'ok': False, 'error': 'No se pudo guardar la asistencia. Intentá nuevamente.'}), 500
    if attendance['status'] == 'unavailable':
        return jsonify({'ok': False, 'error': attendance['message']}), 409
    return jsonify({
        'ok': True,
        'face_detected': True,
        'identified': True,
        'user': {'id': alumno['id'], 'nombre': alumno['nombre'], 'apellido': alumno['apellido']},
        'distance': round(distance, 4),
        'attendance': attendance,
    })

@app.route('/entrada/foto/<int:alumno_id>', methods=['POST'])
@login_required
@rol_requerido('entrada')
def guardar_foto_alumno(alumno_id):
    data = request.get_json(silent=True) or {}
    photo_data = data.get('photo')
    if not isinstance(photo_data, str) or not photo_data.startswith(('data:image/jpeg;base64,', 'data:image/png;base64,')):
        return jsonify({'ok': False, 'error': 'La foto debe ser JPEG o PNG.'}), 400
    if len(photo_data) > 8 * 1024 * 1024:
        return jsonify({'ok': False, 'error': 'La foto supera el tamaño permitido.'}), 413

    conn = get_db()
    alumno = conn.execute(
        "SELECT id, rol, profile_photo, face_consent FROM usuarios WHERE id=?", (alumno_id,)
    ).fetchone()
    if not alumno or alumno['rol'] != 'alumno':
        conn.close()
        return jsonify({'ok': False, 'error': 'Alumno no encontrado.'}), 404
    if alumno['profile_photo']:
        conn.close()
        return jsonify({'ok': False, 'error': 'El alumno ya tiene una foto.'}), 409

    filename = None
    try:
        filename = save_photo(photo_data=photo_data)
        face_encoding = None
        if alumno['face_consent']:
            face_encoding = serialize_encoding(extract_single_encoding_from_file(
                os.path.join(app.config['UPLOAD_FOLDER'], filename)
            ))
        cursor = conn.execute(
            "UPDATE usuarios SET profile_photo=?, face_encoding=? WHERE id=? AND rol='alumno' "
            "AND (profile_photo IS NULL OR profile_photo='') AND face_consent=?",
            (filename, face_encoding, alumno_id, alumno['face_consent'])
        )
        if cursor.rowcount != 1:
            conn.rollback()
            os.remove(os.path.join(app.config['UPLOAD_FOLDER'], filename))
            filename = None
            return jsonify({'ok': False, 'error': 'La foto del alumno cambió. Actualizá la página.'}), 409
        conn.commit()
        return jsonify({'ok': True, 'photo_url': url_for('static', filename='uploads/' + filename)})
    except (ValueError, FaceImageError, FaceRecognitionUnavailable) as error:
        conn.rollback()
        if filename:
            try:
                os.remove(os.path.join(app.config['UPLOAD_FOLDER'], filename))
            except OSError:
                app.logger.exception('No se pudo limpiar una foto no asignada')
        return jsonify({'ok': False, 'error': str(error)}), 400
    except (OSError, sqlite3.Error):
        conn.rollback()
        if filename:
            try:
                os.remove(os.path.join(app.config['UPLOAD_FOLDER'], filename))
            except OSError:
                app.logger.exception('No se pudo limpiar una foto no asignada')
        app.logger.exception('No se pudo guardar la foto del alumno %s', alumno_id)
        return jsonify({'ok': False, 'error': 'No se pudo guardar la foto. Intentá nuevamente.'}), 500
    finally:
        conn.close()

@app.route('/entrada/marcar/<int:alumno_id>', methods=['POST'])
@login_required
@rol_requerido('entrada', 'preceptor')
def marcar_entrada(alumno_id):
    ahora = datetime.now()
    fecha = ahora.date().isoformat()
    hora = ahora.strftime('%H:%M:%S')
    conn = get_db()
    alumno = conn.execute("SELECT id, turno, hora_entrada_oficial, rol, qr_token FROM usuarios WHERE id=?", (alumno_id,)).fetchone()
    if not alumno or alumno['rol'] != 'alumno':
        conn.close()
        flash('El QR no corresponde a un alumno.')
        return redirect(url_for('dashboard'))
    existente = conn.execute("SELECT id FROM asistencias WHERE alumno_id=? AND fecha=?", (alumno_id, fecha)).fetchone()
    if existente:
        conn.close()
        return redirect(url_for('ver_usuario_por_qr', token=alumno['qr_token']))
    inicio = get_official_start(alumno)
    estado = 'tardanza' if inicio and (ahora.hour, ahora.minute) > inicio else 'presente'
    conn.execute("INSERT INTO asistencias (alumno_id, fecha, estado, hora) VALUES (?,?,?,?)",
                 (alumno_id, fecha, estado, hora))
    conn.commit()
    conn.close()
    flash(f'Entrada registrada a las {hora}. Estado: {estado}.')
    panel = 'dashboard_entrada' if session.get('rol') == 'entrada' else 'dashboard_preceptor'
    return redirect(url_for(panel, scan=1))

# --- Asistencia (preceptor) ---
@app.route('/asistencia/registrar', methods=['GET', 'POST'])
@login_required
@rol_requerido('preceptor')
def registrar_asistencia():
    curso = session.get('curso')
    if not curso:
        flash('No tienes un curso asignado.')
        return redirect(url_for('dashboard'))
    conn = get_db()
    alumnos = conn.execute("SELECT id, nombre, apellido, turno, profile_photo FROM usuarios WHERE rol='alumno' AND curso=? ORDER BY nombre",
                           (curso,)).fetchall()
    if request.method == 'POST':
        fecha = request.form['fecha']
        for alumno in alumnos:
            estado = request.form.get(f'estado_{alumno["id"]}')
            if estado:
                # Actualizar o insertar
                existente = conn.execute("SELECT id FROM asistencias WHERE alumno_id=? AND fecha=?",
                                         (alumno['id'], fecha)).fetchone()
                if existente:
                    conn.execute("UPDATE asistencias SET estado=? WHERE id=?",
                                 (estado, existente['id']))
                else:
                    conn.execute("INSERT INTO asistencias (alumno_id, fecha, estado) VALUES (?,?,?)",
                                 (alumno['id'], fecha, estado))
        conn.commit()
        conn.close()
        flash('Asistencia registrada correctamente.')
        return redirect(url_for('ver_asistencia_curso'))
    conn.close()
    hoy = date.today().isoformat()
    return render_template('registrar_asistencia.html', alumnos=alumnos, hoy=hoy)

@app.route('/asistencia/curso', methods=['GET'])
@login_required
@rol_requerido('preceptor')
def ver_asistencia_curso():
    curso = session.get('curso')
    if not curso:
        flash('No tienes un curso asignado.')
        return redirect(url_for('dashboard'))
    fecha_filtro = request.args.get('fecha', date.today().isoformat())
    conn = get_db()
    alumnos = conn.execute("SELECT id, nombre, apellido, turno, profile_photo FROM usuarios WHERE rol='alumno' AND curso=? ORDER BY nombre",
                           (curso,)).fetchall()
    registros = {}
    for alumno in alumnos:
        reg = conn.execute("SELECT estado, hora FROM asistencias WHERE alumno_id=? AND fecha=?",
                           (alumno['id'], fecha_filtro)).fetchone()
        registros[alumno['id']] = dict(reg) if reg else {'estado': 'Sin registro', 'hora': None}
    conn.close()
    return render_template('ver_asistencia.html', alumnos=alumnos, registros=registros,
                           fecha=fecha_filtro, es_preceptor=True)

# --- Alumno ve su asistencia ---
@app.route('/asistencia/mi', methods=['GET'])
@login_required
@rol_requerido('alumno')
def ver_mi_asistencia():
    alumno_id = session['usuario_id']
    mes = request.args.get('mes', datetime.now().strftime('%Y-%m'))
    conn = get_db()
    registros = conn.execute("SELECT fecha, estado FROM asistencias WHERE alumno_id=? AND strftime('%Y-%m', fecha)=? ORDER BY fecha",
                             (alumno_id, mes)).fetchall()
    # Estadísticas
    total = len(registros)
    presentes = sum(1 for r in registros if r['estado'] == 'presente')
    ausentes = sum(1 for r in registros if r['estado'] == 'ausente')
    tardanzas = sum(1 for r in registros if r['estado'] == 'tardanza')
    conn.close()
    return render_template('ver_asistencia.html', registros=registros, mes=mes,
                           total=total, presentes=presentes, ausentes=ausentes,
                           tardanzas=tardanzas, es_alumno=True)

# --- Directivo: reporte por curso ---
@app.route('/asistencia/reporte', methods=['GET'])
@login_required
@rol_requerido('directivo')
def reporte_asistencia():
    curso = request.args.get('curso', '')
    fecha = request.args.get('fecha', date.today().isoformat())
    conn = get_db()
    cursos = [row[0] for row in conn.execute("SELECT DISTINCT curso FROM usuarios WHERE rol='alumno' AND curso IS NOT NULL").fetchall()]
    alumnos = []
    registros = {}
    if curso:
        alumnos = conn.execute("SELECT id, nombre FROM usuarios WHERE rol='alumno' AND curso=? ORDER BY nombre",
                               (curso,)).fetchall()
        for al in alumnos:
            reg = conn.execute("SELECT estado FROM asistencias WHERE alumno_id=? AND fecha=?",
                               (al['id'], fecha)).fetchone()
            registros[al['id']] = reg['estado'] if reg else 'Sin registro'
    conn.close()
    return render_template('reporte_asistencia.html', cursos=cursos, curso_sel=curso,
                           fecha=fecha, alumnos=alumnos, registros=registros)

if __name__ == '__main__':
    host = os.environ.get('FLASK_HOST', '0.0.0.0')
    port = int(os.environ.get('FLASK_PORT', '5000'))
    debug = os.environ.get('FLASK_DEBUG', '0') == '1'
    use_https = os.environ.get('FLASK_HTTPS', '1') == '1'
    cert_file = os.environ.get('FLASK_SSL_CERT')
    key_file = os.environ.get('FLASK_SSL_KEY')
    if bool(cert_file) != bool(key_file):
        raise RuntimeError('Configura FLASK_SSL_CERT y FLASK_SSL_KEY juntos.')
    if cert_file and key_file:
        ssl_context = (cert_file, key_file)
    else:
        ssl_context = 'adhoc' if use_https else None
    app.run(host=host, port=port, debug=debug, ssl_context=ssl_context)